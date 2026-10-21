"""Check preservation of user cells and agreement of native VBA with repository sources."""
from pathlib import Path
from copy import copy
from zipfile import ZipFile
from xml.etree import ElementTree as E
import argparse
import hashlib
import json
import re
import warnings
import subprocess
from decimal import Decimal, localcontext
import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import column_index_from_string

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
parser = argparse.ArgumentParser()
parser.add_argument('--baseline', type=Path, default=HERE / 'Baseline/RC_Section_NDM.xlsm')
parser.add_argument('--candidate', type=Path, required=True)
parser.add_argument('--native', type=Path, required=True)
parser.add_argument('--baseline-native', type=Path, required=True)
args = parser.parse_args()
original = args.baseline
assert hashlib.sha256(original.read_bytes()).hexdigest().upper() == '7B4C769D3813CE0057FFA93026281CD8ED7D3DC23F60159A2655882AB37D2772', 'Baseline mismatch; inspect the supplied original workbook.'

with warnings.catch_warnings():
    warnings.simplefilter('ignore', UserWarning)
    before = openpyxl.load_workbook(original, rich_text=True)
    after = openpyxl.load_workbook(args.candidate, rich_text=True)
assert before.sheetnames == after.sheetnames
setting = next(c for c in after['Config']['A'] if c.value == 'General.ReformatResultsEnabled')
inserted = setting.row
assert after['Config'].cell(inserted, 2).value == 'Yes'
guide = next(c.row for row in after['Справка'] for c in row if isinstance(c.value, str) and c.value.startswith('General.ReformatResultsEnabled —'))
guide_end = max(c.row for row in after['Справка'] for c in row if c.value is not None)

def map_ref(value, current_sheet):
    # Only references into the five shifted Config columns move.
    if value.startswith(('Config!', "'Config'!")):
        current_sheet = 'Config'
    pattern = r"(?:(?P<sheet>'[^']+'|[A-Za-z_А-Яа-я][\w ]*)!)?(?P<col>\$?[A-Z]{1,3})(?P<row>\$?\d+)"
    def part(match):
        sheet = match['sheet'].strip("'") if match['sheet'] else current_sheet
        row = int(match['row'].replace('$',''))
        if sheet == 'Config' and column_index_from_string(match['col'].replace('$','')) <= 5 and row >= inserted:
            return (match[0][:-len(match['row'])] + ('$' if match['row'].startswith('$') else '') + str(row+1))
        return match[0]
    return re.sub(pattern, part, value)

def map_formula(value, sheet):
    if not isinstance(value, str) or not value.startswith('='):
        return value
    tokens = Tokenizer(value).items
    return '=' + ''.join(map_ref(t.value, sheet) if t.type == 'OPERAND' and t.subtype == 'RANGE' else t.value for t in tokens)

def style(cell):
    return tuple(copy(getattr(cell, key)) for key in ['font','fill','border','alignment','protection']) + (cell.number_format,)

checked = 0
validation_cells = 0
for first, second in zip(before, after):
    for (row, col), cell in first._cells.items():
        new_row = row + 1 if first.title == 'Config' and col <= 5 and row >= inserted else row
        if first.title == 'Справка' and guide <= row <= guide_end:
            assert cell.value is None
            assert not cell.comment and not getattr(cell,'hyperlink',None)
            continue
        target = second.cell(new_row, col)
        assert target.value == map_formula(cell.value, first.title), (first.title, cell.coordinate, str(cell.value)[:100], str(target.value)[:100])
        assert style(cell) == style(target), (first.title, cell.coordinate, 'style')
        assert (cell.comment.text if cell.comment else None) == (target.comment.text if target.comment else None)
        if getattr(cell, 'hyperlink', None):
            assert target.hyperlink
            assert cell.hyperlink.target == target.hyperlink.target
            expected = map_ref(cell.hyperlink.location or '', first.title)
            assert (target.hyperlink.location or '').replace("'Config'!", 'Config!') == expected.replace("'Config'!", 'Config!'), (first.title, cell.coordinate, 'link', expected, target.hyperlink.location)
        checked += 1
    assert set(first.column_dimensions) == set(second.column_dimensions)
    for key in first.column_dimensions:
        for attr in ['width','hidden','outlineLevel','collapsed','min','max','bestFit']:
            assert getattr(first.column_dimensions[key],attr) == getattr(second.column_dimensions[key],attr), (first.title,key,attr)
    for key in first.row_dimensions:
        if first.title == 'Справка' and guide <= key <= guide_end:
            continue
        for attr in ['height','hidden','outlineLevel','collapsed','thickTop','thickBot']:
            assert getattr(first.row_dimensions[key],attr) == getattr(second.row_dimensions[key],attr), (first.title,key,attr)
    for validation in first.data_validations.dataValidation:
        for area in validation.sqref.ranges:
            assert area.max_row - area.min_row < 10000, 'Inspect a whole-column validation explicitly.'
            for row in range(area.min_row, area.max_row + 1):
                for col in range(area.min_col, area.max_col + 1):
                    new_row = row + 1 if first.title == 'Config' and col <= 5 and row >= inserted else row
                    address = second.cell(new_row,col).coordinate
                    matches = [v for v in second.data_validations.dataValidation if address in v.sqref]
                    assert len(matches) == 1, (first.title,address,'validation coverage')
                    current = matches[0]
                    for key in ['type','operator','allowBlank','showErrorMessage','errorTitle','error','promptTitle','prompt']:
                        assert getattr(validation,key) == getattr(current,key), (first.title,address,key)
                    for key in ['formula1','formula2']:
                        value = getattr(validation,key)
                        assert getattr(current,key) == map_formula(value,first.title), (first.title,address,key)
                    validation_cells += 1

assert set(before.defined_names) == set(after.defined_names)
for name in before.defined_names:
    expected = map_ref(before.defined_names[name].attr_text, '')
    assert after.defined_names[name].attr_text == expected, (name, expected, after.defined_names[name].attr_text)

drawings = []
with ZipFile(original) as a, ZipFile(args.candidate) as b:
    for name in a.namelist():
        if name.startswith(('xl/drawings/','xl/charts/','xl/media/')):
            assert a.read(name) == b.read(name), ('graphic part',name)
            drawings.append(name)
    ns = {'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    aa = E.fromstring(a.read('xl/workbook.xml')).find('x:calcPr',ns)
    bb = E.fromstring(b.read('xl/workbook.xml')).find('x:calcPr',ns)
    assert aa.get('calcMode','auto') == bb.get('calcMode','auto')

def normalize(text):
    lines = text.replace('\r\n','\n').splitlines()
    text = '\n'.join(line.rstrip() for line in lines if not line.startswith('Attribute ')).strip()
    # VBE changes identifier casing across modules. Strings and comments stay exact.
    def code_part(part):
        def number(match):
            token = match[0]
            # VBE prints this existing Double sentinel with 15 significant digits.
            if token == '1.797693134862315e+308': token = '1.79769313486231e+308'
            suffix = token[-1] if token[-1] in '#%!@&' else ''
            if suffix: token = token[:-1]
            if suffix == '#' and Decimal(token) != Decimal(token).to_integral(): suffix = ''
            with localcontext() as context:
                context.prec = 100
                return str(Decimal(token.replace('d','e')).normalize()) + suffix
        return re.sub(r'(?<![\w.])(?:\d+(?:\.\d*)?|\.\d+)(?:e[+-]?\d+)?[#%!@&]?',number,part.lower())
    return ''.join(part.encode('cp1251','replace').decode('cp1251') if part.startswith("'") else part if part.startswith('"') else code_part(part) for part in re.findall(
        r'"(?:[^"\n]|"")*"|\x27[^\n]*|[^"\x27\n]+|\n',text))

def native(path):
    text = path.read_text(encoding='utf-8-sig')
    pattern = r'(?m)^COMPONENT: ([^\n]+)\nTYPE: [^\n]+\nLINES: \d+\n={100}\n'
    matches = list(re.finditer(pattern, text.replace('\r\n','\n')))
    assert matches
    text = text.replace('\r\n','\n')
    return {m[1].strip(): normalize(text[m.end(): matches[i+1].start()-101 if i+1<len(matches) else len(text)]) for i,m in enumerate(matches)}

actual, old_native = native(args.native), native(args.baseline_native)
sources = {}
for directory in ['src','tests']:
    for file in (ROOT / directory).rglob('*'):
        if file.suffix not in ['.bas','.cls','.frm']: continue
        code = file.read_text(encoding='utf-8-sig')
        match = re.search(r'^Attribute VB_Name = "([^"]+)"',code,re.M)
        assert match, file
        explicit = re.search(r'^Option Explicit(?:\n|$)',code,re.M)
        body = code[explicit.start():] if explicit else '\n'.join(line for line in code.splitlines() if not line.startswith('Attribute '))
        # New comments must be representable in the VBE code page; pre-existing
        # Greek comments were already '?' in the saved book and remain unchanged.
        previous = subprocess.check_output(['git','show','33181b12ab9b8bcb49a72f22fab3a0905819ca61:'+file.relative_to(ROOT).as_posix()]).decode('utf-8-sig') if file.name != 'modResultsFormatting.bas' else ''
        old_comments = {line.split("'",1)[1] for line in previous.splitlines() if "'" in line}
        for line in body.splitlines():
            if "'" not in line: continue
            comment = line.split("'",1)[1]
            if comment.encode('cp1251','replace').decode('cp1251') != comment:
                assert comment in old_comments, ('New unrepresentable comment',file,line)
        sources[match[1]] = normalize(body)
for name, expected in sources.items():
    assert actual.get(name) == expected, ('VBA mismatch',name)
documents = set(actual) - set(sources)
for name in documents:
    assert actual[name] == old_native[name], ('User/document component changed',name)
assert not any(name in actual for name in ['modSpeedTradeoffProbe','modCompactJsonProbe','modResultsFormattingProbe'])
result = {'Passed': True, 'PreservedCells': checked, 'PreservedValidationCells': validation_cells, 'GraphicPartsByteExact': len(drawings),
          'Setting': 'General.ReformatResultsEnabled', 'Default': 'Yes', 'ConfigRow': inserted,
          'HelpRows': [guide,guide_end], 'SourceComponents': len(sources), 'NativeComponents': len(actual),
          'DocumentComponentsUnchanged': sorted(documents), 'WorkbookSHA256': hashlib.sha256(args.candidate.read_bytes()).hexdigest().upper(),
          'AllCodeSHA256': hashlib.sha256(args.native.read_bytes()).hexdigest().upper()}
(HERE / 'Preservation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
