"""Check unchanged user workbook data and actual VBA after the acceptance fixes."""
from pathlib import Path
from copy import copy
from zipfile import ZipFile
from decimal import Decimal, localcontext
import argparse, json, hashlib, re, warnings
import openpyxl

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[2]
parser=argparse.ArgumentParser()
parser.add_argument('--candidate',type=Path,required=True)
parser.add_argument('--native',type=Path,required=True)
args=parser.parse_args()
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest().upper()
baseline=HERE/'Baseline.xlsm'
assert sha(baseline)=='0002E910E28AD95B03362296315839C245AAACAD80011B2ACA158115A69F0845'
with warnings.catch_warnings():
 warnings.simplefilter('ignore',UserWarning)
 before=openpyxl.load_workbook(baseline,rich_text=True)
 after=openpyxl.load_workbook(args.candidate,rich_text=True)
assert before.sheetnames==after.sheetnames
cells=0
for a,b in zip(before,after):
 assert set(a._cells)==set(b._cells),(a.title,'cell inventory')
 for key,c in a._cells.items():
  d=b._cells[key]
  assert c.value==d.value,(a.title,c.coordinate,'value')
  for field in ['font','fill','alignment','protection','number_format']:
   assert copy(getattr(c,field))==copy(getattr(d,field)),(a.title,c.coordinate,field)
  if a.title=='Config' and c.coordinate=='E7':
   assert copy(d.border)==copy(a['E6'].border),'Expected full dotted border'
  else:assert copy(c.border)==copy(d.border),(a.title,c.coordinate,'border')
  assert (c.comment.text if c.comment else None)==(d.comment.text if d.comment else None)
  if getattr(c,'hyperlink',None):
   assert d.hyperlink and c.hyperlink.target==d.hyperlink.target
   assert (c.hyperlink.location or '').replace("'Config'!",'Config!')==(d.hyperlink.location or '').replace("'Config'!",'Config!')
  cells+=1
 assert list(map(str,a.merged_cells.ranges))==list(map(str,b.merged_cells.ranges))
 for key,x in a.column_dimensions.items():
  y=b.column_dimensions[key]
  for field in ['width','hidden','outlineLevel','collapsed','min','max','bestFit']:assert getattr(x,field)==getattr(y,field),(a.title,key,field)
 for key,x in a.row_dimensions.items():
  y=b.row_dimensions[key]
  for field in ['height','hidden','outlineLevel','collapsed','thickTop','thickBot']:assert getattr(x,field)==getattr(y,field),(a.title,key,field)
 assert str(a.data_validations)==str(b.data_validations),(a.title,'validation')
assert set(before.defined_names)==set(after.defined_names)
for name in before.defined_names:assert before.defined_names[name].attr_text==after.defined_names[name].attr_text,name
parts=0
with ZipFile(baseline) as a,ZipFile(args.candidate) as b:
 for name in a.namelist():
  if name.startswith(('xl/drawings/','xl/charts/','xl/media/')):
   assert a.read(name)==b.read(name),name
   parts+=1
# Use the existing explicit VBE normalization (case/literal spelling/codepage).
helper=(ROOT/'docs/regression/ResultsFormatting_2026-10-21/Verify-Candidate.py').read_text(encoding='utf-8')
exec(helper[helper.index('def normalize('):helper.index('actual, old_native = ')])
actual=native(args.native);old=native(HERE/'BaselineVBA.txt')
assert set(actual)==set(old)
changed=[];source_names=set();production=0
for directory in ['src','tests']:
 for file in (ROOT/directory).rglob('*'):
  if file.suffix not in ['.bas','.cls','.frm']:continue
  text=file.read_text(encoding='utf-8-sig')
  name=re.search(r'^Attribute VB_Name = "([^"]+)"',text,re.M)[1]
  source_names.add(name)
  start=re.search(r'^Option Explicit(?:\n|$)',text,re.M)
  expected=normalize(text[start.start():] if start else '\n'.join(x for x in text.splitlines() if not x.startswith('Attribute ')))
  assert expected==actual[name],('source/native',name)
  if actual[name]!=old[name]:changed.append(name)
  if directory=='src':assert actual[name]==old[name],('Production module changed',name);production+=1
for name in set(actual)-source_names:assert actual[name]==old[name],('Document module',name)
assert set(changed)=={'modTestWorkbookInterface','modTestPlotConfig','modTestRecipeOperations','modTestGeometryReview','modTestBatchCalculation'},changed
result={'Passed':True,'PreservedCells':cells,'OnlyCellFormattingChange':'Config!E7 full border','GraphicPartsByteExact':parts,'NativeComponents':len(actual),'SourceComponents':len(source_names),'ProductionComponentsUnchanged':production,'ChangedTestComponents':changed,'WorkbookSHA256':sha(args.candidate),'AllCodeSHA256':sha(args.native)}
(HERE/'Preservation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
